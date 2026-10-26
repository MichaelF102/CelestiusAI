"""Dashboard engine: model, filters, KPIs, figures, interactions, exports, code, templates."""

import ast
import io
import json

import numpy as np
import pandas as pd
import pytest
from src.agent.llm_provider import MockLLMProvider
from src.datascience import dashboard as db
from src.datascience import dashboard_ai as ai
from src.datascience import dashboard_code as dcode
from src.datascience import dashboard_export as dx
from streamlit.testing.v1 import AppTest


@pytest.fixture(scope="module")
def df():
    rng = np.random.default_rng(11)
    n = 360
    out = pd.DataFrame(
        {
            "customer_id": np.arange(n),
            "day": pd.date_range("2024-01-01", periods=n).astype(str),
            "region": rng.choice(["North", "South", "East", "West"], n),
            "country": rng.choice(["France", "Germany", "Spain", "Italy"], n),
            "plan": rng.choice(["basic", "pro", "team"], n),
            "revenue": rng.gamma(2, 60, n).round(2),
            "cost": rng.gamma(2, 30, n).round(2),
            "units": rng.integers(1, 20, n),
            "churn": (rng.random(n) < 0.25).astype(int),
        }
    )
    out.loc[::17, "cost"] = np.nan
    return out


@pytest.fixture(autouse=True)
def _no_show(monkeypatch):
    import plotly.graph_objects as go

    monkeypatch.setattr(go.Figure, "show", lambda self, *a, **k: None)


def T(kind, title="t", **params):
    options = params.pop("options", {})
    return db.Tile(kind=kind, title=title, params=params, options=options)


def S(*tiles, **kw):
    return db.DashboardSpec(tiles=list(tiles), **kw)


# --------------------------------------------------------------------------- model
def test_old_spec_files_still_load():
    old = json.dumps(
        {
            "title": "Old",
            "columns": 2,
            "tiles": [{"id": "a", "kind": "bar", "title": "B", "params": {"x": "region"}}],
            "filters": [{"column": "region", "kind": "categorical", "values": ["N"]}],
        }
    )
    spec = db.spec_from_json(old)
    assert (
        spec.tiles[0].width == 1
        and spec.tiles[0].source == "active"
        and spec.theme.template == "dark"
    )
    assert db.spec_from_json(db.spec_to_json(spec)) == spec
    for bad in ("[]", "nonsense", '{"tiles": [{"kind": "nope", "title": "x"}]}'):
        with pytest.raises(db.DashboardError):
            db.spec_from_json(bad)


def test_sections_keep_first_appearance_order():
    spec = S(T("kpi", "a"), T("kpi", "b"), T("kpi", "c"))
    spec.tiles[0].section = "Z"
    spec.tiles[1].section = "A"
    spec.tiles[2].section = "Z"
    assert spec.sections == ["Z", "A"]


# --------------------------------------------------------------------------- parameters & calculated fields
def test_parameters_resolve_recursively():
    spec = S(
        T("kpi", "k", column="revenue", agg={"$param": "agg"}),
        parameters=[db.Parameter(name="agg", value="sum"), db.Parameter(name="n", value=3)],
    )
    spec.tiles[0].options = {"goal": {"$param": "n"}, "refs": [{"value": {"$param": "n"}}]}
    r = db.resolve_tile(spec.tiles[0], spec)
    assert (
        r.params["agg"] == "sum" and r.options["goal"] == 3 and r.options["refs"][0]["value"] == 3
    )
    with pytest.raises(db.DashboardError, match="Unknown parameter"):
        db.resolve({"$param": "ghost"}, {})


def test_calculated_fields(df):
    fields = [
        db.CalcField(name="profit", expr="revenue - cost"),
        db.CalcField(name="margin", expr="profit / revenue"),
        db.CalcField(name="revenue", expr="units * 2"),
        db.CalcField(name="bad", expr="__import__('os')"),
        db.CalcField(name="ghost", expr="nope + 1"),
    ]
    out, errors = db.apply_calculated(df, fields)
    assert out["profit"].iloc[1] == pytest.approx(df["revenue"].iloc[1] - df["cost"].iloc[1])
    assert "margin" in out.columns
    assert set(errors) == {"revenue", "bad", "ghost"} and "already exists" in errors["revenue"]
    assert "profit" not in df.columns  # input untouched


# --------------------------------------------------------------------------- filters
def test_text_preset_single_and_selective_filters(df):
    assert set(
        db.apply_filters(df, [db.Filter(column="region", kind="text", contains="ORTH")])["region"]
    ) == {"North"}
    last = db.apply_filters(df, [db.Filter(column="day", kind="date", preset="last_30")])
    dates = pd.to_datetime(df["day"])
    assert len(last) == (dates >= dates.max() - pd.Timedelta(days=30)).sum()
    year = db.apply_filters(df, [db.Filter(column="day", kind="date", preset="this_year")])
    assert pd.to_datetime(year["day"]).dt.year.eq(pd.to_datetime(df["day"]).max().year).all()
    assert len(db.apply_filters(df, [db.Filter(column="day", kind="date", preset="all")])) == len(
        df
    )
    with pytest.raises(db.DashboardError):
        db.apply_filters(df, [db.Filter(column="day", kind="date", preset="last_year")])
    f = db.Filter(column="region", kind="categorical", values=["North"], tiles=["only_me"])
    assert len(db.apply_filters(df, [f], tile_id="other")) == len(df)
    assert set(db.apply_filters(df, [f], tile_id="only_me")["region"]) == {"North"}


def test_tile_data_layers_calculated_global_local_and_cross(df):
    tables = {"active": df}
    spec = S(
        T("kpi", "k", agg="count"),
        filters=[db.Filter(column="region", kind="categorical", values=["North", "South"])],
        calculated=[db.CalcField(name="profit", expr="revenue - cost")],
    )
    tile = spec.tiles[0]
    tile.local_filters = [db.Filter(column="plan", kind="categorical", values=["pro"])]
    out = db.tile_data(
        tables, spec, tile, [db.Filter(column="country", kind="categorical", values=["France"])]
    )
    exp = df[df.region.isin(["North", "South"]) & (df.plan == "pro") & (df.country == "France")]
    assert len(out) == len(exp) and "profit" in out.columns
    tile.apply_global_filters = False
    assert len(db.tile_data(tables, spec, tile)) == (df.plan == "pro").sum()
    tile.source = "ghost"
    with pytest.raises(db.DashboardError, match="not available"):
        db.tile_data(tables, spec, tile)


# --------------------------------------------------------------------------- KPIs
def test_kpi_previous_period_target_goal_threshold_spark(df):
    dates = pd.to_datetime(df["day"])
    end = dates.max()
    cur = df[(dates > end - pd.Timedelta(days=30)) & (dates <= end)]["revenue"].sum()
    prev = df[(dates > end - pd.Timedelta(days=60)) & (dates <= end - pd.Timedelta(days=30))][
        "revenue"
    ].sum()
    tile = T(
        "kpi",
        "Revenue",
        column="revenue",
        agg="sum",
        options={
            "compare": {"kind": "previous_period", "date_col": "day", "days": 30},
            "sparkline": {"date_col": "day", "freq": "W"},
            "thresholds": [{"op": "<", "value": 1e9, "color": "red"}],
            "prefix": "$",
        },
    )
    card = db.kpi_card(df, tile)
    assert card["value"] == pytest.approx(cur) and card["delta"] == pytest.approx(cur - prev)
    assert card["delta_text"] == f"{(cur - prev) / prev:+.1%}" and card["direction"] in (
        "up",
        "down",
    )
    assert len(card["spark"]) > 40 and card["color"] == "red" and card["text"].startswith("$")
    total = df["revenue"].sum()
    tgt = db.kpi_card(
        df,
        T(
            "kpi",
            "t",
            column="revenue",
            agg="sum",
            options={"compare": {"kind": "target", "target": total * 2}},
        ),
    )
    assert tgt["progress"] == pytest.approx(0.5) and "vs target" in tgt["delta_text"]
    goal = db.kpi_card(df, T("kpi", "g", column="revenue", agg="sum", options={"goal": total / 4}))
    assert goal["progress"] == 1.0
    rate = db.kpi_card(df, T("kpi", "r", column="churn", agg="mean", options={"format": "percent"}))
    assert rate["text"] == f"{df.churn.mean() * 100:.1f}%"
    assert db.kpi_card(df.head(0), T("kpi", "e", column="revenue", agg="sum"))["text"] == "-"


def test_value_formats():
    assert (
        db.format_value(0.256, "percent") == "25.6%"
        and db.format_value(1234.5, "currency") == "$1,234"
    )
    assert (
        db.format_value(1234.5, "integer") == "1,234"
        or db.format_value(1234.5, "integer") == "1,235"
    )
    assert (
        db.format_value(3.14159, "decimal", decimals=3) == "3.142" and db.format_value(None) == "-"
    )
    assert (
        db.format_value(2_500_000) == "2.50M"
        and db.format_value(15_000, prefix="€", suffix="!") == "€15.0K!"
    )


# --------------------------------------------------------------------------- frames & figures
TILES = {
    "bar": T("bar", x="region", y="revenue", agg="sum", color="plan"),
    "area": T("area", x="plan", y="revenue", agg="mean"),
    "line": T("line", x="day", y="revenue", agg="sum", freq="W"),
    "combo": T("combo", x="day", y="revenue", y2="cost", agg="sum", freq="MS"),
    "scatter": T("scatter", x="units", y="revenue", color="plan", options={"trendline": True}),
    "hist": T("hist", column="revenue", bins=20, color="plan"),
    "pie": T("pie", names="region", values="revenue", agg="sum"),
    "box": T("box", y="revenue", x="plan"),
    "heatmap": T("heatmap", x="plan", y="region", value="revenue", agg="mean"),
    "funnel": T("funnel", names="plan", values="revenue", agg="sum"),
    "waterfall": T("waterfall", x="region", y="revenue", agg="sum"),
    "treemap": T("treemap", path=["region", "plan"], values="revenue", agg="sum"),
    "map": T("map", names="country", values="revenue", agg="sum"),
    "gauge": T("gauge", column="revenue", agg="sum", min=0, max=100000),
}


@pytest.mark.parametrize("kind", sorted(TILES))
@pytest.mark.parametrize("theme", ["dark", "light"])
def test_every_tile_kind_builds_a_figure(kind, theme, df):
    t = TILES[kind]
    fig = db.build_figure(df, t, theme=db.Theme(template=theme, palette="colorblind", font_size=14))
    assert fig is not None and len(fig.data) >= 1 and fig.layout.height == 340
    assert fig.layout.template.layout.paper_bgcolor is not None
    assert fig.layout.font.size == 14


def test_latlon_map_text_table_and_unknown(df):
    geo = df.assign(lat=np.linspace(-60, 60, len(df)), lon=np.linspace(-170, 170, len(df)))
    assert (
        len(
            db.build_figure(geo, T("map", mode="latlon", lat="lat", lon="lon", size="revenue")).data
        )
        == 1
    )
    assert (
        db.build_figure(df, T("text", markdown="hi")) is None
        and db.build_figure(df, T("table")) is None
    )
    assert db.tile_frame(df, T("table", columns=["region"], rows=5)).shape == (5, 1)
    with pytest.raises(db.DashboardError):
        db.tile_frame(df, T("bar", x="ghost"))
    with pytest.raises(db.DashboardError, match="filters"):
        db.tile_frame(df.head(0), TILES["bar"])


def test_bar_sort_top_n_other_and_stacking(df):
    t = T(
        "bar",
        x="country",
        y="revenue",
        agg="sum",
        options={"sort": "desc", "top_n": 2, "other_group": True},
    )
    frame = db.tile_frame(df, t)
    assert len(frame) == 3 and frame["country"].iloc[-1] == "Other"
    assert frame["revenue"].sum() == pytest.approx(df["revenue"].sum())
    assert frame["revenue"].iloc[0] >= frame["revenue"].iloc[1]
    asc = db.tile_frame(df, T("bar", x="country", y="revenue", agg="sum", options={"sort": "asc"}))
    assert asc["revenue"].is_monotonic_increasing
    mean_top = db.tile_frame(
        df,
        T("bar", x="country", y="revenue", agg="mean", options={"top_n": 2, "other_group": True}),
    )
    assert "Other" not in set(
        mean_top["country"]
    )  # "Other" only makes sense for additive aggregates
    stacked = db.build_figure(
        df,
        T(
            "bar",
            x="region",
            y="revenue",
            agg="sum",
            color="plan",
            options={"stacked": "percent", "data_labels": True},
        ),
    )
    assert stacked.layout.barmode == "relative" and stacked.layout.barnorm == "percent"
    horizontal = db.build_figure(
        df, T("bar", x="region", y="revenue", agg="sum", options={"orientation": "h"})
    )
    assert horizontal.data[0].orientation == "h"


def test_values_match_pandas(df):
    bar = db.tile_frame(df, T("bar", x="region", y="revenue", agg="sum")).set_index("region")[
        "revenue"
    ]
    pd.testing.assert_series_equal(
        bar.sort_index(), df.groupby("region")["revenue"].sum().sort_index(), check_names=False
    )
    heat = db.tile_frame(df, TILES["heatmap"])
    assert heat.loc["North", "pro"] == pytest.approx(
        df[(df.region == "North") & (df.plan == "pro")]["revenue"].mean()
    )
    pie = db.tile_frame(
        df,
        T(
            "pie",
            names="region",
            values="revenue",
            agg="sum",
        ),
    )
    assert pie["revenue"].sum() == pytest.approx(df["revenue"].sum())
    line = db.tile_frame(df, TILES["line"])
    assert line["revenue"].sum() == pytest.approx(df["revenue"].sum())
    combo = db.build_figure(df, TILES["combo"])
    assert len(combo.data) == 2 and combo.data[1].yaxis == "y2"
    wf = db.build_figure(df, TILES["waterfall"])
    assert list(wf.data[0].measure)[-1] == "total" and len(wf.data[0].x) == df.region.nunique() + 1
    funnel = db.tile_frame(df, TILES["funnel"])
    assert funnel["revenue"].is_monotonic_decreasing


def test_display_options_reach_the_figure(df):
    t = T(
        "line",
        x="day",
        y="revenue",
        agg="sum",
        freq="MS",
        options={
            "log_y": True,
            "x_title": "When",
            "y_title": "Money",
            "ref_lines": [
                {"kind": "mean", "label": "avg"},
                {"kind": "value", "value": 500, "label": "target"},
            ],
        },
    )
    t.height = 500
    fig = db.build_figure(df, t)
    assert (
        fig.layout.yaxis.type == "log"
        and fig.layout.xaxis.title.text == "When"
        and fig.layout.yaxis.title.text == "Money"
    )
    assert fig.layout.height == 500 and len(fig.layout.shapes) == 2
    assert fig.layout.shapes[1].y0 == 500


# --------------------------------------------------------------------------- captions & validation
def test_captions_are_computed_from_the_data(df):
    top = df.groupby("region")["revenue"].sum().idxmax()
    assert top in db.caption_for(df, T("bar", x="region", y="revenue", agg="sum"))
    assert "correlation" in db.caption_for(df, TILES["scatter"])
    assert f"{len(df):,} of {len(df):,}" in db.caption_for(df, T("table", rows=500))
    assert (
        db.caption_for(df, T("text", markdown="x")) == ""
        and db.caption_for(df, T("bar", x="ghost")) == ""
    )
    assert "Revenue" in db.caption_for(df, T("kpi", "Revenue", column="revenue", agg="sum"))


def test_validate_spec_reports_problems(df):
    spec = S(
        T("bar", "ok", x="region", y="revenue", agg="sum"),
        T("bar", "gone", x="ghost"),
        T("kpi", "away", agg="count"),
        filters=[db.Filter(column="nope", kind="text", contains="x")],
        calculated=[db.CalcField(name="bad", expr="zzz + 1")],
    )
    spec.tiles[2].source = "elsewhere"
    problems = db.validate_spec(spec, {"active": df})
    text = " | ".join(problems)
    assert "'gone': missing column(s) ghost" in text and "'away': table 'elsewhere'" in text
    assert "Filter on 'nope'" in text and "Calculated field 'bad'" in text and "'ok'" not in text


# --------------------------------------------------------------------------- interactions
def test_cross_filter_from_selection():
    bar = T("bar", x="region", y="revenue")
    sel = {"selection": {"points": [{"x": "North"}, {"x": "South"}, {"x": "North"}]}}
    f = db.selection_to_filter(bar, sel)
    assert f.column == "region" and f.values == ["North", "South"]
    h = T("bar", x="region", y="revenue", options={"orientation": "h"})
    assert db.selection_to_filter(h, {"selection": {"points": [{"y": "East"}]}}).values == ["East"]
    assert (
        db.selection_to_filter(
            T("pie", names="plan", values="revenue"), {"selection": {"points": [{"label": "pro"}]}}
        ).column
        == "plan"
    )
    assert (
        db.selection_to_filter(bar, {"selection": {"points": []}}) is None
        and db.selection_to_filter(bar, None) is None
    )


def test_drill_down_and_up(df):
    tile = T("bar", x="region", y="revenue", agg="sum", drill=["region", "country", "plan"])
    state = db.DrillState()
    t0, f0 = db.drill_apply(tile, state)
    assert t0.params["x"] == "region" and f0 == []
    state = db.drill_down(tile, state, "North")
    t1, f1 = db.drill_apply(tile, state)
    assert t1.params["x"] == "country" and f1[0].values == ["North"]
    frame = db.tile_frame(db.apply_filters(df, f1), t1)
    assert frame["revenue"].sum() == pytest.approx(df[df.region == "North"]["revenue"].sum())
    state = db.drill_down(tile, state, "France")
    t2, f2 = db.drill_apply(tile, state)
    assert t2.params["x"] == "plan" and [f.column for f in f2] == ["region", "country"]
    assert db.drill_down(tile, state, "pro") == state  # already at the deepest level
    up = db.drill_up(state)
    assert up.level == 1 and list(up.picked) == ["region"]
    assert db.drill_up(db.DrillState()).level == 0


# --------------------------------------------------------------------------- layout & history
def test_layout_rows_move_duplicate():
    tiles = [db.Tile(kind="kpi", title=str(i), width=w) for i, w in enumerate([2, 1, 3, 1, 1, 1])]
    rows = db.layout_rows(tiles, 3)
    assert [[t.title for t in r] for r in rows] == [["0", "1"], ["2"], ["3", "4", "5"]]
    assert [[t.title for t in r] for r in db.layout_rows(tiles, 2)][2] == [
        "2"
    ]  # span capped to the grid
    spec = S(*tiles)
    moved = db.move_tile(spec, tiles[0].id, 1)
    assert [t.title for t in moved.tiles][:2] == ["1", "0"] and [t.title for t in spec.tiles][
        0
    ] == "0"
    assert db.move_tile(spec, tiles[0].id, -1).tiles[0].title == "0"
    dup = db.duplicate_tile(spec, tiles[1].id)
    assert (
        dup.tiles[2].title == "1 (copy)" and dup.tiles[2].id != tiles[1].id and len(dup.tiles) == 7
    )


def test_history_undo_redo():
    h = db.History(limit=3)
    s0 = S(T("kpi", "a"))
    s1 = S(T("kpi", "a"), T("kpi", "b"))
    h.push(s0)
    assert h.undo(s1).tiles[0].title == "a" and len(h.undo_stack) == 0
    redone = h.redo(s0)
    assert len(redone.tiles) == 2 and h.redo(redone) is None
    for i in range(6):
        h.push(S(*[T("kpi", str(k)) for k in range(i + 1)]))
    assert len(h.undo_stack) == 3 and not h.redo_stack
    h.push(S(*[T("kpi", str(k)) for k in range(6)]))  # same snapshot: no duplicate
    assert len(h.undo_stack) == 3
    assert db.History().undo(s0) is None


# --------------------------------------------------------------------------- exports
def rich_spec():
    spec = S(
        T(
            "text",
            "Notes",
            markdown="### Hello <script>alert(1)</script>\n- **bold** item\n[ok](https://example.com) [bad](javascript:alert(1))",
        ),
        T(
            "kpi",
            "Revenue",
            column="revenue",
            agg="sum",
            options={
                "compare": {"kind": "previous_period", "date_col": "day", "days": 30},
                "sparkline": {"date_col": "day", "freq": "W"},
                "goal": 1e5,
            },
        ),
        TILES["bar"].model_copy(update={"title": "By <region>", "width": 2}),
        T(
            "table",
            "Rows",
            columns=["region", "revenue"],
            rows=5,
            options={
                "formats": {"revenue": "currency"},
                "highlight": [{"column": "revenue", "op": ">", "value": 100, "color": "red"}],
            },
        ),
        title="Sales <b>",
        description="A & B",
        theme=db.Theme(template="light", accent="#ff0000", font_size=13),
    )
    spec.tiles[1].section = spec.tiles[2].section = "Overview"
    spec.filters = [db.Filter(column="region", kind="categorical", values=["North"])]
    return spec


def test_html_export_is_safe_and_complete(df):
    page = dx.export_html(df, rich_spec())
    assert "<script>alert" not in page and "&lt;script&gt;" in page
    assert '<a href="javascript' not in page  # the unsafe link stays inert text
    assert (
        "<title>Sales &lt;b&gt;</title>" in page
        and "A &amp; B" in page
        and "By &lt;region&gt;" in page
    )
    assert "<h2>Overview</h2>" in page and "grid-column: span 2" in page and "#ff0000" in page
    assert (
        "<polyline" in page
        and "▲" in page + "▲"
        and 'role="img"' in page
        and "<figcaption>" in page
    )
    assert "<strong>bold</strong>" in page and '<a href="https://example.com"' in page
    assert "$" in page and "background:#ef4444" in page  # currency format + highlight rule
    assert f"{(df.region == 'North').sum():,} of {len(df):,} rows" in page
    assert page.count("cdn.plot.ly") + page.count("plotly.min.js") >= 1


def test_html_export_survives_percent_signs_and_broken_tiles(df):
    spec = S(T("bar", "broken", x="ghost"), T("text", "t", markdown="100% sure %s %d"))
    page = dx.export_html({"active": df}, spec)
    assert "100% sure %s %d" in page and "Column(s) not found" in page


def test_markdown_converter():
    out = dx.md_to_html("# T\n*it* **b**\n- a\n- b\n\n<img src=x onerror=1>")
    assert (
        "<h3>T</h3>" in out
        and "<em>it</em>" in out
        and out.count("<li>") == 2
        and "<img" not in out
    )


def test_xlsx_export_has_one_sheet_per_data_tile(df):
    from openpyxl import load_workbook

    spec = rich_spec()
    data = dx.export_xlsx(df, spec)
    wb = load_workbook(io.BytesIO(data))
    assert (
        wb.sheetnames[0] == "Summary"
        and "By _region_" in wb.sheetnames
        or any(n.startswith("By") for n in wb.sheetnames)
    )
    summary = pd.read_excel(io.BytesIO(data), sheet_name="Summary")
    assert {"tile", "type", "rows", "sheet"} <= set(summary.columns) and "Notes" not in set(
        summary["tile"]
    )
    bar_sheet = next(n for n in wb.sheetnames if n.startswith("By"))
    frame = pd.read_excel(io.BytesIO(data), sheet_name=bar_sheet)
    assert set(frame["region"]) == {"North"}  # global filter applied


def test_pdf_export_with_and_without_images(df, monkeypatch):
    assert dx.figure_png(object()) is None  # kaleido missing / failing -> None, not an exception
    pdf, charts = dx.export_pdf(df, rich_spec())
    assert pdf.startswith(b"%PDF") and charts is False
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    buf = io.BytesIO()
    plt.figure(figsize=(2, 1)).savefig(buf, format="png")
    monkeypatch.setattr(dx, "figure_png", lambda fig, *a, **k: buf.getvalue())
    pdf2, charts2 = dx.export_pdf(df, rich_spec())
    assert charts2 is True and pdf2.startswith(b"%PDF")


# --------------------------------------------------------------------------- code generation
def test_snippet_runs_for_a_rich_dashboard(df, tmp_path):
    csv = tmp_path / "d.csv"
    df.to_csv(csv, index=False)
    spec = S(
        *[
            TILES[k]
            for k in (
                "bar",
                "area",
                "line",
                "combo",
                "scatter",
                "hist",
                "pie",
                "box",
                "heatmap",
                "funnel",
                "waterfall",
                "treemap",
                "map",
            )
        ],
        T(
            "kpi",
            "Revenue",
            column="revenue",
            agg="sum",
            options={
                "compare": {"kind": "previous_period", "date_col": "day", "days": 30},
                "format": "currency",
            },
        ),
        T("gauge", "G", column="revenue", agg="sum"),
        T("table", "Rows", columns=["region", "revenue"], rows=5),
        T("text", "Note", markdown="hello"),
        filters=[
            db.Filter(column="region", kind="categorical", values=["North", "South"]),
            db.Filter(column="units", kind="numeric", low=3, high=15),
            db.Filter(column="day", kind="date", preset="last_90"),
            db.Filter(column="plan", kind="text", contains="a"),
        ],
        calculated=[db.CalcField(name="profit", expr="revenue - cost")],
        parameters=[db.Parameter(name="agg", value="sum")],
    )
    spec.tiles[0].options = {
        "sort": "desc",
        "top_n": 3,
        "other_group": True,
        "log_y": True,
        "ref_lines": [{"kind": "value", "value": 100, "label": "t"}],
        "x_title": "X",
    }
    spec.tiles[0].section = "S1"
    code = dcode.dashboard_snippet(spec, filename=str(csv))
    ast.parse(code)
    exec(compile(code, "<dash>", "exec"), {"__name__": "x"})  # noqa: S102


@pytest.mark.parametrize("kind", sorted(set(TILES) - {"gauge"}))
def test_generated_frames_match_the_live_frames(kind, df, tmp_path):
    csv = tmp_path / "d.csv"
    df.to_csv(csv, index=False)
    tile = TILES[kind].model_copy(deep=True)
    spec = S(tile, filters=[db.Filter(column="plan", kind="categorical", values=["pro", "basic"])])
    ns: dict = {"__name__": "x"}
    exec(compile(dcode.dashboard_snippet(spec, filename=str(csv)), "<d>", "exec"), ns)  # noqa: S102
    live = db.tile_frame(db.tile_data({"active": df}, spec, tile), tile)
    got = ns["frame"]
    if kind == "heatmap":
        pd.testing.assert_frame_equal(
            got.sort_index().sort_index(axis=1),
            live.sort_index().sort_index(axis=1),
            check_dtype=False,
            check_names=False,
        )
    else:
        key = list(live.columns)
        a = got[key].sort_values(key).reset_index(drop=True)
        b = live[key].sort_values(key).reset_index(drop=True)
        pd.testing.assert_frame_equal(a, b, check_dtype=False)


def test_generated_kpi_matches_live_card(df, tmp_path):
    csv = tmp_path / "d.csv"
    df.to_csv(csv, index=False)
    tile = T(
        "kpi",
        "Revenue",
        column="revenue",
        agg="sum",
        options={
            "compare": {"kind": "previous_period", "date_col": "day", "days": 30},
            "format": "currency",
        },
    )
    ns: dict = {"__name__": "x"}
    exec(compile(dcode.dashboard_snippet(S(tile), filename=str(csv)), "<d>", "exec"), ns)  # noqa: S102
    card = db.kpi_card(df, tile)
    assert ns["kpi_text"] == card["text"] and ns["kpi_delta"] == card["delta_text"]


def test_streamlit_app_runs_and_filters_react(df, tmp_path):
    csv = tmp_path / "d.csv"
    df.to_csv(csv, index=False)
    spec = S(
        T("kpi", "Rows", agg="count"),
        T("kpi", "Revenue", column="revenue", agg="sum", options={"format": "currency"}),
        TILES["bar"],
        TILES["pie"],
        T("table", "Rows", columns=["region", "revenue"], rows=5),
        T("text", "Note", markdown="**hi**"),
        title="Gen App",
        filters=[
            db.Filter(column="region", kind="categorical", values=[]),
            db.Filter(column="units", kind="numeric"),
            db.Filter(column="day", kind="date", preset="all"),
            db.Filter(column="plan", kind="text"),
        ],
        calculated=[db.CalcField(name="profit", expr="revenue - cost")],
    )
    code = dcode.streamlit_app_code(spec, str(csv))
    ast.parse(code)
    at = AppTest.from_string(code, default_timeout=60).run()
    assert not at.exception, [e.value for e in at.exception]
    metrics = {m.label: m.value for m in at.metric}
    assert metrics["Rows"] == f"{len(df):,}" and metrics["Revenue"].startswith("$")
    assert len(at.dataframe) == 1 and any("hi" in m.value for m in at.markdown)
    at.sidebar.multiselect[0].set_value(["North"]).run()
    assert not at.exception
    assert {m.label: m.value for m in at.metric}["Rows"] == f"{(df.region == 'North').sum():,}"
    at.sidebar.text_input[0].set_value("pro").run()
    assert {m.label: m.value for m in at.metric}[
        "Rows"
    ] == f"{((df.region == 'North') & df.plan.str.contains('pro')).sum():,}"


# --------------------------------------------------------------------------- templates & AI
@pytest.mark.parametrize("name", ["auto", "sales", "customer", "operations"])
def test_templates_produce_renderable_specs(name, df):
    spec = ai.template_spec(name, df)
    assert len(spec.tiles) >= 5 and spec.sections
    tables = {"active": df}
    for tile in spec.tiles:
        data = db.tile_data(tables, spec, tile)
        if tile.kind == "kpi":
            assert db.kpi_card(data, tile)["text"] != ""
        elif tile.kind not in ("text", "table"):
            assert db.build_figure(data, tile, theme=spec.theme) is not None
        else:
            db.tile_frame(data, tile)
    assert db.validate_spec(spec, tables) == []


def test_quality_template_uses_its_own_table(df):
    spec = ai.template_spec("quality", df)
    qt = ai.quality_table(df.assign(extra=np.nan))
    assert qt.loc[qt.column == "extra", "missing_pct"].iloc[0] == 100.0
    tables = {"active": df, "quality_summary": qt}
    assert db.validate_spec(spec, tables) == []
    assert db.validate_spec(spec, {"active": df})  # table missing -> reported
    for tile in spec.tiles:
        db.tile_frame(db.tile_data(tables, spec, tile), tile)


def test_suggest_adds_findings_and_template_errors(df):
    class F:
        def __init__(self, title, severity):
            self.title, self.severity = title, severity

    spec = ai.suggest_spec(df, findings=[F("Revenue is skewed", "high"), F("minor", "low")])
    assert spec.tiles[0].kind == "text" and "Revenue is skewed" in spec.tiles[0].params["markdown"]
    with pytest.raises(db.DashboardError):
        ai.template_spec("nonsense", df)
    with pytest.raises(db.DashboardError, match="too few"):
        ai.template_spec(
            "sales",
            pd.DataFrame({"a": [f"u{i}" for i in range(60)], "b": [f"v{i}" for i in range(60)]}),
        )


def test_roles_detection(df):
    r = ai.ColumnRoles(df)
    assert r.measures[0] in ("revenue", "cost") and "churn" in r.binary and "customer_id" in r.ids
    assert "region" in r.categories and "day" in r.dates


def test_plain_language_tile_is_validated(df):
    good = MockLLMProvider(
        predefined_response='Sure:\n{"kind": "bar", "title": "Revenue by region", "params": {"x": "region", "y": "revenue", "agg": "sum"}}'
    )
    tile = ai.nl_to_tile(good, "revenue by region", df)
    assert tile.kind == "bar" and db.tile_frame(df, tile)["revenue"].sum() == pytest.approx(
        df.revenue.sum()
    )
    for text, msg in [
        ('{"kind": "bar", "title": "x", "params": {"x": "ghost"}}', "unknown column"),
        ("no json here", "valid tile"),
        ('{"kind": "bubble", "title": "x"}', "valid tile"),
    ]:
        with pytest.raises(db.DashboardError, match=msg):
            ai.nl_to_tile(MockLLMProvider(predefined_response=text), "x", df)
    prompt = ai.build_tile_prompt("hello", df)
    assert "revenue" in prompt and "Request: hello" in prompt


def test_pandas_failures_become_friendly_errors(df):
    with pytest.raises(db.DashboardError, match="different columns"):
        db.tile_frame(df, T("line", x="revenue", y="revenue", agg="sum"))
    with pytest.raises(db.DashboardError, match="Could not"):
        db.tile_frame(df, T("bar", x="region", y="region", agg="sum"))  # cannot sum text
    with pytest.raises(db.DashboardError):
        db.build_figure(df, T("scatter", x="units", y="units"))
